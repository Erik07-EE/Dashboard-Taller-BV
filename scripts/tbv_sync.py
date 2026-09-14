# -*- coding: utf-8 -*-
"""
Sync del Dashboard Taller BV — Electroestrada.

Uso:
    python scripts/tbv_sync.py preview --dump <ruta-del-volcado>
    python scripts/tbv_sync.py aplicar
    python scripts/tbv_sync.py verificar

`preview`  decodifica el XLSX, extrae produccion + pedidos, compara contra el
           dashboard e imprime el PREVIEW. No modifica nada.
`aplicar`  escribe los cambios en el HTML (requiere un `preview` previo).
`verificar` compara el HTML actual contra el estado anterior, campo por campo.

El volcado es el .txt que deja el conector de Google Drive cuando la respuesta
de download_file_content excede el limite de tokens (ver CLAUDE.md).
"""
import argparse
import base64
import json
import os
import re
import shutil
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HTML = os.path.join(RAIZ, 'Dashboard_Pedidos.html')
CACHE = os.path.join(os.environ.get('TEMP', os.path.expanduser('~')), 'tbv-sync')

BS = chr(92)
AR = timezone(timedelta(hours=-3))
MESES_ES = {1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio',
            7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre',
            12: 'Diciembre'}
GA_HOJA = [('Inducidos', 'OF-IND'), ('Rotores', 'OF-ROT'), ('Estatores', 'OF-EST')]
GA_ORD = {'Inducidos': 0, 'Rotores': 1, 'Estatores': 2}

K_SEM = ['num', 'dias', 'obj', 'prod']
K_COD = ['codigo', 'categoria', 'cantidad']
K_OF = ['ga', 'mes_key', 'mes_label', 'obj_mensual', 'dias_habiles', 'prod_total']
# Orden canonico. Cambiarlo marca como modificados registros que no cambiaron.
K_PED = ['n', 'ga', 'codigo', 'condicion', 'cliente', 'reclamo', 'observacion',
         'fecha_ingreso', 'dias_objetivo', 'entrega_sol', 'fecha_entrega_real',
         'dias_prod', 'demora']


def ruta_larga(p):
    r"""Prefijo \\?\ para superar el limite de 260 caracteres de Windows."""
    p = os.path.abspath(p)
    if os.name == 'nt' and len(p) > 250 and not p.startswith(BS * 2):
        return BS * 2 + '?' + BS + p
    return p


# --------------------------------------------------------------------------
# Parser de array-literal JS (claves sin comillas, strings con comilla simple)
# --------------------------------------------------------------------------
def js_literal_to_json(s):
    out, i, n = [], 0, len(s)
    while i < n:
        c = s[i]
        if c == "'":
            j, buf = i + 1, []
            while j < n:
                if s[j] == BS:
                    nxt = s[j + 1] if j + 1 < n else ''
                    buf.append("'" if nxt == "'" else s[j] + nxt)
                    j += 2
                    continue
                if s[j] == "'":
                    break
                buf.append(s[j])
                j += 1
            out.append(json.dumps(''.join(buf)))
            i = j + 1
            continue
        if c == '"':
            j, buf = i + 1, ['"']
            while j < n:
                if s[j] == BS:
                    buf.append(s[j] + (s[j + 1] if j + 1 < n else ''))
                    j += 2
                    continue
                buf.append(s[j])
                if s[j] == '"':
                    break
                j += 1
            out.append(''.join(buf))
            i = j + 1
            continue
        m = re.match(r'[A-Za-z_$][\w$]*', s[i:])
        if m:
            word = m.group(0)
            if re.match(r'\s*:', s[i + len(word):]) or word not in ('true', 'false', 'null'):
                out.append(json.dumps(word))
            else:
                out.append(word)
            i += len(word)
            continue
        out.append(c)
        i += 1
    return ''.join(out)


def extraer_array(html, nombre):
    """Devuelve (inicio, fin, texto) del array-literal `const <nombre>=[...]`."""
    marca = 'const ' + nombre + '=['
    start = html.index(marca) + len(marca) - 1
    depth, i, in_str, q = 0, start, False, ''
    while i < len(html):
        ch = html[i]
        if in_str:
            if ch == BS:
                i += 2
                continue
            if ch == q:
                in_str = False
        elif ch in ("'", '"'):
            in_str, q = True, ch
        elif ch == '[':
            depth += 1
        elif ch == ']':
            depth -= 1
            if depth == 0:
                return start, i + 1, html[start:i + 1]
        i += 1
    raise ValueError('no cierra ' + nombre)


def leer_html():
    # newline='' preserva los finales de linea LF del archivo.
    with open(HTML, encoding='utf-8', newline='') as f:
        return f.read()


def leer_dashboard(html=None):
    html = html or leer_html()
    res = {}
    for nombre in ('OF_DATA', 'PEDIDOS'):
        _, _, txt = extraer_array(html, nombre)
        res[nombre] = json.loads(js_literal_to_json(txt))
    return res


# --------------------------------------------------------------------------
# Serializacion JS-literal
# --------------------------------------------------------------------------
def jsval(v):
    if v is None:
        return 'null'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        f = float(v)
        return str(int(f)) if f == int(f) else repr(round(f, 4))
    s = str(v).replace(BS, BS * 2).replace("'", BS + "'")
    return "'" + s + "'"


def jsobj(d, claves):
    return '{' + ','.join(k + ':' + jsval(d.get(k)) for k in claves) + '}'


def ser_of(e):
    base = ','.join(k + ':' + jsval(e.get(k)) for k in K_OF)
    sems = '[' + ','.join(
        jsobj(s, [k for k in K_SEM if k != 'obj' or s.get('obj') is not None])
        for s in e.get('semanas', [])) + ']'
    cods = '[' + ','.join(jsobj(c, K_COD) for c in e.get('codigos', [])) + ']'
    return '{' + base + ',semanas:' + sems + ',codigos:' + cods + '}'


# --------------------------------------------------------------------------
# Extraccion del XLSX
# --------------------------------------------------------------------------
def decodificar(dump, destino):
    with open(ruta_larga(dump), encoding='utf-8') as f:
        d = json.load(f)
    raw = base64.b64decode(d['content'])
    if raw[:2] != b'PK':
        raise SystemExit('ERROR: el contenido no es un XLSX (falta header PK)')
    with open(destino, 'wb') as f:
        f.write(raw)
    return len(raw), d.get('title')


def extraer_ga(wb, sheet_name, year, month):
    ws = wb[sheet_name]
    mes_nombre = MESES_ES[month]
    maxc = ws.max_column
    row1 = {c: ws.cell(1, c).value for c in range(1, maxc + 1)}

    lab_c = None
    for c in range(1, maxc + 1):
        v = row1[c]
        if isinstance(v, str) and v.strip().rstrip(':').lower() == ('objetivo ' + mes_nombre).lower():
            lab_c = c
            break
    obj_mensual = row1.get(lab_c + 1) if lab_c else None

    # Semanales rotulados: entre el 'Objetivo' anterior y este. NO hacia adelante,
    # porque agarraria los del mes siguiente.
    obj_sems = {}
    if lab_c:
        prev = 0
        for c in range(lab_c - 1, 0, -1):
            v = row1.get(c)
            if isinstance(v, str) and v.strip().lower().startswith('objetivo'):
                prev = c
                break
        for c in range(prev + 1, lab_c):
            v = row1.get(c)
            if isinstance(v, str) and 'Obj. S.' in v:
                try:
                    num = int(v.split('.')[-1].strip())
                    val = row1.get(c + 1)
                    if isinstance(val, (int, float)):
                        obj_sems[num] = round(val)
                except ValueError:
                    pass

    cols = []
    for c in range(1, maxc + 1):
        v = ws.cell(2, c).value
        if isinstance(v, datetime) and v.year == year and v.month == month:
            if sheet_name == 'OF-EST' and v.month == 4 and v.day == 30:
                continue  # artifact conocido en Estatores
            cols.append((c, v))
    if not cols:
        return None

    total, codigos = 0, []
    for row in range(3, ws.max_row + 1):
        cod = ws.cell(row, 1).value
        if not cod or not isinstance(cod, str):
            continue
        cod = cod.strip()
        if cod in ('-', ''):
            continue  # filas de total
        p = sum(ws.cell(row, c).value or 0 for c, _ in cols
                if isinstance(ws.cell(row, c).value, (int, float)))
        if p > 0:
            total += p
            cat = ws.cell(row, 2).value
            codigos.append({'codigo': cod.replace(',', '.'),
                            'categoria': str(cat).strip() if cat else '',
                            'cantidad': int(p)})

    obj_m = round(obj_mensual) if isinstance(obj_mensual, (int, float)) else 0
    base = {'mes_key': '%04d-%02d' % (year, month),
            'mes_label': '%s %d' % (mes_nombre, year),
            'obj_mensual': obj_m, 'dias_habiles': len(cols)}
    if obj_m == 0 and total == 0:
        base.update({'prod_total': 0, 'semanas': [], 'codigos': []})
        return base

    sem_map = defaultdict(list)
    for c, dt in sorted(cols, key=lambda x: x[1]):
        sem_map[dt.isocalendar()[1]].append((c, dt))

    semanas = []
    for i, (_wk, ccols) in enumerate(sorted(sem_map.items())):
        wn = i + 1
        objw = obj_sems.get(wn)
        if objw is None:
            # Semana de 1 dia sin rotulo: valor suelto en fila 1.
            for c, _ in ccols:
                v, pv = row1.get(c), row1.get(c - 1)
                if isinstance(v, (int, float)) and not (isinstance(pv, str) and 'Prod' in pv):
                    objw = round(v)
                    break
            if objw is None:
                objw = 0
        prod_s = 0
        for row in range(3, ws.max_row + 1):
            cod = ws.cell(row, 1).value
            if not cod or not isinstance(cod, str) or cod.strip() in ('-', ''):
                continue
            prod_s += sum(ws.cell(row, c).value or 0 for c, _ in ccols
                          if isinstance(ws.cell(row, c).value, (int, float)))
        semanas.append({'num': wn, 'dias': len(ccols), 'obj': objw, 'prod': int(prod_s)})

    if obj_m:
        suma = sum(x['obj'] for x in semanas)
        if suma != obj_m:
            sin_rotulo = [x for x in semanas if x['num'] not in obj_sems]
            if len(sin_rotulo) == 1:
                sin_rotulo[0]['obj'] = obj_m - sum(
                    x['obj'] for x in semanas if x['num'] in obj_sems)
            else:
                print('  [AVISO] %s %s: suma semanales (%d) != mensual (%d) - revisar sheet'
                      % (sheet_name, mes_nombre, suma, obj_m))

    base.update({'prod_total': int(total), 'semanas': semanas, 'codigos': codigos})
    return base


def fmt_date(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, str) and v not in ('-', ''):
        return v
    return None


def extraer_pedidos(wb):
    ws = wb['Urgentes']
    pedidos = {}
    for row in range(3, ws.max_row + 1):
        n, ga = ws.cell(row, 1).value, ws.cell(row, 3).value
        if not isinstance(n, (int, float)) or not isinstance(ga, str):
            continue
        n = int(n)
        dpr = ws.cell(row, 16).value
        cod = ws.cell(row, 4).value
        obj = ws.cell(row, 7).value
        dem = ws.cell(row, 17).value
        pedidos[n] = {
            'n': n, 'ga': str(ga).strip(),
            'codigo': str(cod).strip().replace(',', '.') if cod else None,
            'condicion': ws.cell(row, 5).value, 'cliente': ws.cell(row, 6).value,
            'fecha_ingreso': fmt_date(ws.cell(row, 2).value),
            'dias_objetivo': int(obj) if isinstance(obj, (int, float)) else None,
            'entrega_sol': fmt_date(ws.cell(row, 8).value),
            'reclamo': ws.cell(row, 12).value if ws.cell(row, 12).value not in (None, '-') else None,
            'observacion': ws.cell(row, 13).value if ws.cell(row, 13).value not in (None, '-') else None,
            'fecha_entrega_real': fmt_date(ws.cell(row, 14).value),
            # dias_prod negativo = error de formula en el sheet
            'dias_prod': int(dpr) if isinstance(dpr, (int, float)) and dpr >= 0 else None,
            'demora': round(float(dem), 4) if isinstance(dem, (int, float)) else None,
        }
    return pedidos


def extraer_todo(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    hoy = datetime.now(AR)
    meses = [(hoy.year, hoy.month)]
    # Meses futuros ya armados en el sheet. Los historicos quedan congelados.
    for m in range(hoy.month + 1, 13):
        if extraer_ga(wb, 'OF-IND', hoy.year, m) is not None:
            meses.append((hoy.year, m))
    bloques = {}
    for (y, m) in meses:
        for ga, hoja in GA_HOJA:
            r = extraer_ga(wb, hoja, y, m)
            if r:
                bloques.setdefault(r['mes_key'], {})[ga] = r
    return {'bloques': bloques, 'pedidos': extraer_pedidos(wb),
            'meses': ['%04d-%02d' % t for t in meses],
            'fecha_ar': hoy.strftime('%Y-%m-%d %H:%M')}


# --------------------------------------------------------------------------
# Comparacion
# --------------------------------------------------------------------------
def es_cambio(campo, v_dash, v_sheet):
    """No sobreescribir datos existentes con vacios del sheet."""
    if v_sheet is None or v_dash == v_sheet:
        return False
    if campo == 'demora' and v_dash is not None:
        try:
            if round(float(v_dash), 2) == round(float(v_sheet), 2):
                return False
        except (TypeError, ValueError):
            pass
    if campo == 'cliente' and str(v_dash or '').lower() == str(v_sheet or '').lower():
        return False
    if campo == 'observacion' and str(v_dash or '').strip() == str(v_sheet or '').strip():
        return False
    return True


def cmd_preview(args):
    os.makedirs(CACHE, exist_ok=True)
    xlsx = os.path.join(CACHE, 'urgentes.xlsx')
    tam, titulo = decodificar(args.dump, xlsx)
    print('XLSX: %s  %d bytes  (header PK OK)' % (titulo, tam))

    sheet = extraer_todo(xlsx)
    dash = leer_dashboard()
    with open(os.path.join(CACHE, 'sheet.json'), 'w', encoding='utf-8') as f:
        json.dump(sheet, f, ensure_ascii=False, default=str)
    with open(os.path.join(CACHE, 'dash.json'), 'w', encoding='utf-8') as f:
        json.dump(dash, f, ensure_ascii=False)

    bloques = sheet['bloques']
    ped_sheet = {int(k): v for k, v in sheet['pedidos'].items()}
    ped_dash = {p['n']: p for p in dash['PEDIDOS']}
    of_dash = {(e['mes_key'], e['ga']): e for e in dash['OF_DATA']}

    print('=' * 72)
    print('PREVIEW - Cambios detectados  [%s]' % sheet['fecha_ar'])
    print('Meses a sincronizar: %s' % ', '.join(sheet['meses']))
    print('=' * 72)

    print('\nOBJETIVOS MENSUALES')
    print('  %-10s %10s %8s  %s' % ('GA', 'Dashboard', 'Sheet', 'Estado'))
    for mk in sorted(bloques):
        for ga in ('Inducidos', 'Rotores', 'Estatores'):
            b = bloques[mk].get(ga)
            if not b:
                continue
            d = of_dash.get((mk, ga))
            od = d['obj_mensual'] if d else None
            print('  %-10s %10s %8s  %s' % (ga, od, b['obj_mensual'],
                                            'OK' if od == b['obj_mensual'] else 'DIFF'))

    print('\nPRODUCCION')
    for mk in sorted(bloques):
        for ga in ('Inducidos', 'Rotores', 'Estatores'):
            b = bloques[mk].get(ga)
            if not b:
                continue
            d = of_dash.get((mk, ga))
            pd_ = d['prod_total'] if d else 0
            sem_d = {s['num']: s.get('prod') for s in (d.get('semanas', []) if d else [])}
            difs = ['S%d: %s->%s' % (s['num'], sem_d.get(s['num']), s['prod'])
                    for s in b['semanas'] if sem_d.get(s['num']) != s['prod']]
            estado = ('%s -> %s' % (pd_, b['prod_total'])) if pd_ != b['prod_total'] \
                else ('%s (sin cambio)' % pd_)
            print('  %s %-10s %-22s %s' % (mk, ga, estado, '  '.join(difs)))
            print('  %s            codigos: %d -> %d'
                  % (' ' * len(mk), len(d.get('codigos', [])) if d else 0, len(b['codigos'])))

    nuevos = sorted(n for n in ped_sheet if n not in ped_dash)
    print('\nPEDIDOS NUEVOS: n=%d' % len(nuevos))
    for n in nuevos:
        p = ped_sheet[n]
        print('  #%-4d [%-10s] %-14s %-16s sol %s'
              % (n, str(p['ga'])[:10], str(p['codigo'])[:14],
                 str(p['condicion'] or '')[:16], p['entrega_sol'] or '-'))

    actualizados = []
    for n, p in ped_sheet.items():
        if n not in ped_dash:
            continue
        d = ped_dash[n]
        difs = [(c, d.get(c), p.get(c)) for c in K_PED
                if c != 'n' and es_cambio(c, d.get(c), p.get(c))]
        if difs:
            actualizados.append((n, d.get('codigo'), difs))

    print('\nPEDIDOS ACTUALIZADOS: n=%d' % len(actualizados))
    for n, cod, difs in sorted(actualizados):
        print('  #%-4d [%-12s] %s' % (n, str(cod)[:12],
              ' | '.join('%s: %s -> %s' % d for d in difs)))

    print('\nTOTALES: pedidos %d -> %d' % (len(ped_dash), len(ped_sheet)))
    print('=' * 72)
    print('Revisar y confirmar. Para aplicar: python scripts/tbv_sync.py aplicar')


def cmd_aplicar(args):
    sheet_p = os.path.join(CACHE, 'sheet.json')
    if not os.path.exists(sheet_p):
        raise SystemExit('ERROR: falta el preview. Corre primero: tbv_sync.py preview --dump ...')
    with open(sheet_p, encoding='utf-8') as f:
        sheet = json.load(f)
    bloques = sheet['bloques']
    ped_sheet = {int(k): v for k, v in sheet['pedidos'].items()}

    backup = os.path.join(CACHE, 'Dashboard_Pedidos.BACKUP.html')
    shutil.copy(HTML, backup)
    html = leer_html()

    # --- OF_DATA ---
    a, b, txt = extraer_array(html, 'OF_DATA')
    of = json.loads(js_literal_to_json(txt))
    idx = {(e['mes_key'], e['ga']): e for e in of}
    tocados = []
    for mk in bloques:
        for ga, nuevo in bloques[mk].items():
            campos = {k: nuevo[k] for k in
                      ('mes_label', 'obj_mensual', 'dias_habiles', 'prod_total',
                       'semanas', 'codigos')}
            if (mk, ga) in idx:
                idx[(mk, ga)].update(campos)
            else:
                e = {'ga': ga, 'mes_key': mk}
                e.update(campos)
                of.append(e)
            tocados.append((mk, ga))
    of.sort(key=lambda e: (e['mes_key'], GA_ORD.get(e['ga'], 9)))
    html = html[:a] + '[\n' + ',\n'.join('  ' + ser_of(e) for e in of) + '\n]' + html[b:]
    print('OF_DATA: %d entradas, actualizadas %s' % (len(of), sorted(set(tocados))))

    # --- PEDIDOS ---
    a, b, txt = extraer_array(html, 'PEDIDOS')
    ped = json.loads(js_literal_to_json(txt))
    previos = {p['n'] for p in ped}
    by_n = {p['n']: p for p in ped}
    n_upd = n_campos = 0
    for n, ps in ped_sheet.items():
        if n in by_n:
            d = by_n[n]
            cambios = [c for c in K_PED if c != 'n' and es_cambio(c, d.get(c), ps.get(c))]
            if cambios:
                n_upd += 1
                n_campos += len(cambios)
                for c in cambios:
                    d[c] = ps.get(c)
        else:
            by_n[n] = {k: ps.get(k) for k in K_PED}
    final = [by_n[n] for n in sorted(by_n)]
    html = html[:a] + '[\n' + ',\n'.join('  ' + jsobj(p, K_PED) for p in final) + '\n]' + html[b:]
    print('PEDIDOS: %d -> %d (nuevos %d, actualizados %d en %d campos)'
          % (len(ped), len(final), len(set(by_n) - previos), n_upd, n_campos))

    # --- badge (re.sub: la linea no tiene punto y coma) ---
    ahora = datetime.now(AR)
    dias = ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']
    fecha = '%s %s' % (dias[ahora.weekday()], ahora.strftime('%d/%m/%Y %H:%M'))
    html, nsub = re.subn(r"const _ULTIMO_SYNC='[^']*'",
                         "const _ULTIMO_SYNC='" + fecha + "'", html)
    if nsub != 1:
        raise SystemExit('ERROR: el badge tuvo %d reemplazos (esperaba 1)' % nsub)
    print("Badge -> '%s'" % fecha)

    # newline='' preserva LF; sin esto Windows reescribe todo el archivo como CRLF
    # y el diff pasa de ~80 lineas a ~1500.
    with open(HTML, 'w', encoding='utf-8', newline='') as f:
        f.write(html)
    print('HTML escrito. Backup: %s' % backup)
    print('Ahora: node scripts/validar.js  y luego  python scripts/tbv_sync.py verificar')


def cmd_verificar(args):
    dash_p = os.path.join(CACHE, 'dash.json')
    if not os.path.exists(dash_p):
        raise SystemExit('ERROR: falta dash.json (corre preview antes de aplicar)')
    with open(dash_p, encoding='utf-8') as f:
        viejo = json.load(f)
    nuevo = leer_dashboard()

    print('OF_DATA: %d -> %d' % (len(viejo['OF_DATA']), len(nuevo['OF_DATA'])))
    print('PEDIDOS: %d -> %d' % (len(viejo['PEDIDOS']), len(nuevo['PEDIDOS'])))

    ov = {(e['mes_key'], e['ga']): e for e in viejo['OF_DATA']}
    on = {(e['mes_key'], e['ga']): e for e in nuevo['OF_DATA']}
    hoy_key = datetime.now(AR).strftime('%Y-%m')
    print('\n--- OF_DATA: bloques con diferencias ---')
    for k in sorted(set(ov) | set(on)):
        if ov.get(k) != on.get(k):
            campos = sorted(c for c in set(list(ov.get(k) or {}) + list(on.get(k) or {}))
                            if (ov.get(k) or {}).get(c) != (on.get(k) or {}).get(c))
            print('  %s %s -> %s' % (k[0], k[1], campos))
    congelados = [k for k in ov if k[0] < hoy_key and ov[k] != on.get(k)]
    print('  MESES CONGELADOS TOCADOS: %s'
          % (congelados if congelados else 'NINGUNO (correcto)'))

    pv = {p['n']: p for p in viejo['PEDIDOS']}
    pn = {p['n']: p for p in nuevo['PEDIDOS']}
    print('\n--- PEDIDOS: diferencias de VALOR (ignora orden de claves) ---')
    cambiados = 0
    for n in sorted(set(pv) & set(pn)):
        difs = sorted(c for c in set(list(pv[n]) + list(pn[n])) if pv[n].get(c) != pn[n].get(c))
        if difs:
            cambiados += 1
            print('  #%d: %s' % (n, ', '.join(
                '%s %s -> %s' % (c, pv[n].get(c), pn[n].get(c)) for c in difs)))
    print('  con cambios de valor: %d | agregados: %d | eliminados: %d'
          % (cambiados, len(set(pn) - set(pv)), len(set(pv) - set(pn))))

    ns = sorted(pn)
    ok = ns == list(range(1, ns[-1] + 1))
    print('\n--- INTEGRIDAD ---')
    print('  n consecutivos 1..%d: %s' % (ns[-1], ok))
    for e in nuevo['OF_DATA']:
        if e['mes_key'] >= hoy_key:
            so = sum(x['obj'] for x in e['semanas'])
            sp = sum(x['prod'] for x in e['semanas'])
            print('  %-10s Sum(obj)=%s vs %s [%s] | Sum(prod)=%s vs %s [%s]'
                  % (e['ga'], so, e['obj_mensual'], 'OK' if so == e['obj_mensual'] else 'ERROR',
                     sp, e['prod_total'], 'OK' if sp == e['prod_total'] else 'ERROR'))
    m = re.search(r"const _ULTIMO_SYNC='([^']*)'", leer_html())
    print('  badge: %s' % (m.group(1) if m else 'NO ENCONTRADO'))


def main():
    ap = argparse.ArgumentParser(description='Sync del Dashboard Taller BV')
    sub = ap.add_subparsers(dest='cmd', required=True)
    p = sub.add_parser('preview', help='extrae y compara, sin modificar nada')
    p.add_argument('--dump', required=True, help='ruta del volcado de download_file_content')
    p.set_defaults(func=cmd_preview)
    p = sub.add_parser('aplicar', help='aplica los cambios al HTML')
    p.set_defaults(func=cmd_aplicar)
    p = sub.add_parser('verificar', help='compara el HTML nuevo contra el anterior')
    p.set_defaults(func=cmd_verificar)
    args = ap.parse_args()
    args.func(args)


if __name__ == '__main__':
    main()
